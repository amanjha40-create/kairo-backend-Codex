"""Secure, deterministic CSV/XLSX parsing for roster previews."""

from __future__ import annotations

import csv
import io
from datetime import date, datetime
from decimal import Decimal
from pathlib import Path
from typing import Any
from zipfile import BadZipFile, ZipFile

from openpyxl import load_workbook

from app.exceptions import ValidationAppError
from app.organization_roster_import.constants import (
    CSV_MIME_TYPES,
    MAX_CELL_LENGTH,
    MAX_COLUMNS,
    MAX_DATA_ROWS,
    MAX_FILENAME_LENGTH,
    MAX_UPLOAD_BYTES,
    MAX_XLSX_ARCHIVE_ENTRIES,
    MAX_XLSX_UNCOMPRESSED_BYTES,
    SUPPORTED_EXTENSIONS,
    XLSX_MIME_TYPES,
)
from app.organization_roster_import.headers import build_source_columns
from app.organization_roster_import.types import ParsedRosterFile, ParsedSourceRow, RowIssue

_OLE_COMPOUND_SIGNATURE = bytes.fromhex("D0CF11E0A1B11AE1")


def parse_roster_file(
    *,
    filename: str,
    content_type: str | None,
    content: bytes,
) -> tuple[str, ParsedRosterFile]:
    """Validate an upload and parse it without evaluating spreadsheet formulas."""

    safe_name, source_format = validate_upload(
        filename=filename,
        content_type=content_type,
        content=content,
    )
    if source_format == "csv":
        return safe_name, parse_csv(content)
    return safe_name, parse_xlsx(content)


def validate_upload(
    *,
    filename: str,
    content_type: str | None,
    content: bytes,
) -> tuple[str, str]:
    original = filename.strip()
    if not original or len(original) > MAX_FILENAME_LENGTH:
        raise ValidationAppError("Filename is missing or too long", code="invalid_filename")
    if "\x00" in original or "/" in original or "\\" in original or Path(original).name != original:
        raise ValidationAppError("Filename must not contain a path", code="unsafe_filename")
    extension = Path(original).suffix.casefold()
    if extension not in SUPPORTED_EXTENSIONS:
        raise ValidationAppError("Only CSV and XLSX files are supported", code="unsupported_file")
    if not content:
        raise ValidationAppError("Roster file is empty", code="empty_file")
    if len(content) > MAX_UPLOAD_BYTES:
        raise ValidationAppError("Roster file exceeds the upload limit", code="file_too_large")

    primary_mime = (content_type or "").split(";", 1)[0].strip().casefold()
    allowed_mimes = CSV_MIME_TYPES if extension == ".csv" else XLSX_MIME_TYPES
    if primary_mime and primary_mime not in allowed_mimes:
        raise ValidationAppError(
            "File content type does not match its extension", code="unsupported_mime_type"
        )
    return original, extension.lstrip(".")


def parse_csv(content: bytes) -> ParsedRosterFile:
    try:
        decoded = content.decode("utf-8-sig")
    except UnicodeDecodeError as exc:
        raise ValidationAppError(
            "CSV must use UTF-8 encoding", code="invalid_csv_encoding"
        ) from exc

    reader = csv.reader(io.StringIO(decoded, newline=""), strict=True)
    try:
        headers = next(reader)
        if not any(_has_value(value) for value in headers):
            raise ValidationAppError("CSV requires a header row", code="missing_header")
        _enforce_column_limit(len(headers))
        _enforce_cells(headers)
        columns = build_source_columns(headers)
        parsed_rows: list[ParsedSourceRow] = []
        previous_line_number = reader.line_num
        for values in reader:
            row_number = previous_line_number + 1
            previous_line_number = reader.line_num
            if len(parsed_rows) >= MAX_DATA_ROWS:
                raise ValidationAppError("Roster exceeds the row limit", code="too_many_rows")
            if len(values) > MAX_COLUMNS:
                raise ValidationAppError("Roster exceeds the column limit", code="too_many_columns")
            _enforce_cells(values)
            blank = not any(_has_value(value) for value in values)
            issues: tuple[RowIssue, ...] = ()
            if not blank and len(values) != len(columns):
                issues = (
                    RowIssue(
                        code="malformed_row",
                        field=None,
                        message="Row has a different number of cells than the header",
                        row_number=row_number,
                    ),
                )
            padded = [*values, *([""] * max(0, len(columns) - len(values)))]
            raw_values = {
                column.original: _json_value(padded[index]) for index, column in enumerate(columns)
            }
            parsed_rows.append(
                ParsedSourceRow(
                    row_number=row_number,
                    raw_values=raw_values,
                    parser_issues=issues,
                    skipped=blank,
                )
            )
    except StopIteration:
        raise ValidationAppError("CSV requires a header row", code="missing_header") from None
    except csv.Error as exc:
        raise ValidationAppError("CSV is malformed", code="malformed_csv") from exc
    _remove_trailing_blank_rows(parsed_rows)
    _require_usable_rows(parsed_rows)
    return ParsedRosterFile(source_format="csv", columns=columns, rows=tuple(parsed_rows))


def parse_xlsx(content: bytes) -> ParsedRosterFile:
    if content.startswith(_OLE_COMPOUND_SIGNATURE):
        raise ValidationAppError(
            "Encrypted or legacy Excel workbooks are not supported",
            code="encrypted_workbook",
        )
    _validate_xlsx_archive(content)
    try:
        workbook = load_workbook(
            io.BytesIO(content),
            read_only=True,
            data_only=False,
            keep_links=False,
        )
    except Exception as exc:
        raise ValidationAppError(
            "XLSX workbook is corrupt or encrypted", code="invalid_xlsx"
        ) from exc

    try:
        non_empty_sheets = [sheet for sheet in workbook.worksheets if _worksheet_has_values(sheet)]
        if not non_empty_sheets:
            raise ValidationAppError("Workbook has no usable rows", code="empty_workbook")
        sheet = non_empty_sheets[0]
        if sheet.max_column > MAX_COLUMNS:
            raise ValidationAppError("Roster exceeds the column limit", code="too_many_columns")

        physical_rows = list(sheet.iter_rows(values_only=True))
        header_offset = next(
            index
            for index, values in enumerate(physical_rows)
            if any(_has_value(v) for v in values)
        )
        header_values = tuple(_json_value(value) for value in physical_rows[header_offset])
        _enforce_cells(header_values)
        columns = build_source_columns(header_values)
        _enforce_column_limit(len(columns))

        parsed_rows: list[ParsedSourceRow] = []
        for zero_index, values in enumerate(
            physical_rows[header_offset + 1 :], start=header_offset + 1
        ):
            if len(parsed_rows) >= MAX_DATA_ROWS:
                raise ValidationAppError("Roster exceeds the row limit", code="too_many_rows")
            trimmed = tuple(values[: len(columns)])
            _enforce_cells(trimmed)
            blank = not any(_has_value(value) for value in trimmed)
            raw_values = {
                column.original: _json_value(trimmed[index] if index < len(trimmed) else None)
                for index, column in enumerate(columns)
            }
            parsed_rows.append(
                ParsedSourceRow(
                    row_number=zero_index + 1,
                    raw_values=raw_values,
                    skipped=blank,
                )
            )
        _remove_trailing_blank_rows(parsed_rows)
        _require_usable_rows(parsed_rows)
        warnings: tuple[str, ...] = ()
        if len(non_empty_sheets) > 1:
            warnings = ("Additional non-empty worksheets were not imported",)
        return ParsedRosterFile(
            source_format="xlsx",
            columns=columns,
            rows=tuple(parsed_rows),
            sheet_name=sheet.title,
            warnings=warnings,
        )
    except ValidationAppError:
        raise
    except Exception as exc:
        raise ValidationAppError(
            "XLSX workbook is corrupt or encrypted", code="invalid_xlsx"
        ) from exc
    finally:
        workbook.close()


def _worksheet_has_values(sheet: Any) -> bool:
    if sheet.max_row > MAX_DATA_ROWS + 1 or sheet.max_column > MAX_COLUMNS:
        raise ValidationAppError(
            "Workbook dimensions exceed roster limits", code="workbook_too_large"
        )
    return any(_has_value(cell) for row in sheet.iter_rows(values_only=True) for cell in row)


def _validate_xlsx_archive(content: bytes) -> None:
    try:
        with ZipFile(io.BytesIO(content)) as archive:
            entries = archive.infolist()
            if len(entries) > MAX_XLSX_ARCHIVE_ENTRIES:
                raise ValidationAppError(
                    "XLSX archive contains too many entries", code="workbook_too_large"
                )
            if sum(entry.file_size for entry in entries) > MAX_XLSX_UNCOMPRESSED_BYTES:
                raise ValidationAppError(
                    "XLSX archive expands beyond the safety limit", code="workbook_too_large"
                )
            if any(entry.flag_bits & 0x1 for entry in entries):
                raise ValidationAppError(
                    "Encrypted workbooks are not supported", code="encrypted_workbook"
                )
            if "xl/workbook.xml" not in archive.namelist():
                raise ValidationAppError("XLSX workbook is corrupt", code="invalid_xlsx")
    except ValidationAppError:
        raise
    except (BadZipFile, OSError, ValueError) as exc:
        raise ValidationAppError("XLSX workbook is corrupt", code="invalid_xlsx") from exc


def _json_value(value: Any) -> Any:
    if value is None:
        return None
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, Decimal):
        return str(value)
    return value


def _has_value(value: Any) -> bool:
    return value is not None and (not isinstance(value, str) or bool(value.strip()))


def _enforce_column_limit(count: int) -> None:
    if count > MAX_COLUMNS:
        raise ValidationAppError("Roster exceeds the column limit", code="too_many_columns")


def _enforce_cells(values: tuple[Any, ...] | list[Any]) -> None:
    for value in values:
        if value is not None and len(str(value)) > MAX_CELL_LENGTH:
            raise ValidationAppError("Roster contains an oversized cell", code="cell_too_long")


def _remove_trailing_blank_rows(rows: list[ParsedSourceRow]) -> None:
    while rows and rows[-1].skipped:
        rows.pop()


def _require_usable_rows(rows: list[ParsedSourceRow]) -> None:
    if not any(not row.skipped for row in rows):
        raise ValidationAppError("Roster has no usable data rows", code="no_usable_rows")
